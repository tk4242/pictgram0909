"""Final Hard Gates を1コマンドで全項目検証する（設計書§32）。

設計書§31のとおり、これはLLMの自主判断ではなく決定論的に走らせるゲート。
1項目でもBLOCKなら READY_FOR_HUMAN_REVIEW へ進ませない。
"""
import json
import re
import sys
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent

n = lambda s: len(unicodedata.normalize("NFC", re.sub(r"\*", "", s.strip())))
results = []


def gate(rid, label, ok, actual, expect, *, stub=False, review=False):
    """review=True の項目は工程を止めない代わりに、人間レビューで必ず提示する。

    設計書§32のHard Gate一覧に無い項目を勝手にBLOCK扱いにすると、
    ゲートの意味（止めるべきものを止める）が曖昧になるため重大度を分ける。
    """
    results.append((rid, label, ok, actual, expect, stub, review))


# ── 記事本文（T/M/B/F/G体系。設計書§13-14）──
lines = [l for l in (HERE / "draft.txt").read_text(encoding="utf-8").splitlines() if l.strip()]
g = lambda p: [l[len(p):] for l in lines if l.startswith(p)]
T, M, H, F, X, G = g("@T@"), g("@M@")[0], g("@H@"), g("@F@"), g("@X@"), g("@G@")
pub, mem = sum(map(n, F)), sum(map(n, X))
tot = pub + mem
KW = ["ベトナム", "イオンモール", "ABCマート", "小売", "出店", "多店舗", "専門店", "5号店", "5店舗"]

gate("T01", "タイトル文字数", all(26 <= n(t) <= 28 for t in T), str([n(t) for t in T]), "26-28字")
gate("T02", "「ベトナム」含有", all("ベトナム" in t for t in T), f"{sum('ベトナム' in t for t in T)}/{len(T)}", "全案")
gate("T03", "タイトル案数", len(T) >= 3, f"{len(T)}案", "3案以上")
gate("M01", "メタディスクリプション", 80 <= n(M) <= 90, f"{n(M)}字", "80-90字")
gate("B01", "本文文字数", 1275 <= tot <= 1725, f"{tot}字", "1,275-1,725字")
gate("B02", "H2見出し数", len(H) >= 2, f"{len(H)}本", "2本以上")
gate("B03", "H2キーワード", not [h for h in H if not any(k in h for k in KW)],
     f"未充足{len([h for h in H if not any(k in h for k in KW)])}本", "全H2")
gate("F01", "図表枚数", len(G) == 2, f"{len(G)}枚", "2枚")
gate("G01", "公開比率", 0.60 <= pub / tot <= 0.70, f"{pub/tot*100:.1f}%（{pub}/{tot}字）", "60-70%")

# ── 会員限定ショートコード（設計書§3）──
cms = (HERE / "cms_output.txt").read_text(encoding="utf-8")
order_ok = (cms.count('[content_control logged_in="true"]') == 1
            and cms.count("[/content_control]") == 1
            and cms.count("[member_cta_buttons]") == 1
            and cms.index("[/content_control]") < cms.index("[member_cta_buttons]"))
gate("SC01", "ショートコード構造", order_ok, "開始1/終了1/CTA1・順序OK" if order_ok else "不正", "各1回・正順")

# ── 図表pptx（設計書§14・§26）──
from pptx import Presentation
prs = Presentation(HERE / "figures.pptx")
km_lens, title_nl, srcs, all_text = [], [], [], []
for slide in prs.slides:
    texts = {sh.name: sh.text_frame.text.strip()
             for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()}
    all_text.append(" ".join(texts.values()))
    for name, t in texts.items():
        if "プレースホルダー 1" in name:
            km_lens.append(n(t))
        if "ボックス 4" in name or "ボックス 5" in name:
            title_nl.append("\n" in t)
        if "プレースホルダー 3" in name:
            srcs.append(t)

gate("F02", "図表キーメッセージ", all(l <= 28 for l in km_lens), str(km_lens), "28字以内")
gate("F03a", "図表タイトル改行なし", not any(title_nl), f"改行{sum(title_nl)}件", "0件（一次判定）")
gate("CIT01", "引用表記", all("InfoBank作成" in s for s in srcs), f"{len(srcs)}件中{sum('InfoBank作成' in s for s in srcs)}件", "全図表")

# ── フォント（設計書§26・E2Eで判明した欠陥）──
z = zipfile.ZipFile(HERE / "figures.pptx")
fonts = Counter()
for name in z.namelist():
    if (name.startswith("ppt/slides/slide") or name.startswith("ppt/charts/chart")) and name.endswith(".xml"):
        fonts += Counter(re.findall(r'typeface="([^"]+)"', z.read(name).decode("utf-8", "ignore")))
bad_fonts = {k: v for k, v in fonts.items() if k != "Meiryo UI"}
gate("FONT01", "Meiryo UI明示指定", bool(fonts.get("Meiryo UI")) and not bad_fonts,
     f"Meiryo UI {fonts.get('Meiryo UI', 0)}箇所" + (f" / 他{bad_fonts}" if bad_fonts else ""), "Meiryo UIのみ")

# ── ブランド表記（設計書§28・確定事項3）──
joined = " ".join(all_text) + " " + cms + " " + "".join(lines)
ng = [b for b in ("InfoBase", "Infobank") if b in joined]
gate("BRAND01", "ブランド表記統一", not ng, f"検出: {ng}" if ng else "InfoBankのみ", "InfoBankのみ")

# ── サムネイル（確定事項2・設計書§22）──
sys.path.insert(0, str(HERE))
from validate_thumbnail import check as thumb_check
import io, contextlib
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    thumb_ok = thumb_check(str(HERE / "thumbnail_typeC.png"),
                           str(REPO / "assets/infobank_logo/infobank_logo_white.png"))
meta = json.loads((HERE / "thumbnail_meta.json").read_text(encoding="utf-8"))
gate("TH01", "サムネイル1枚・構成", thumb_ok, f"TYPE {meta['pattern']}・機械QA{'合格' if thumb_ok else '不合格'}", "1枚・型準拠")
gate("TH02", "ロゴは実ファイル", "抽出した実ロゴ" in meta["logo"]["source"], meta["logo"]["file"], "偽ロゴ禁止")

# ── 外部接続依存（テストモードでスタブ。設計書§32）──
snap = json.loads((HERE / "rules_snapshot.json").read_text(encoding="utf-8"))
gate("RULE01", "ルールhash", bool(snap.get("sha256")), snap["sha256"][:16] + "...", "hash記録", stub=True)
gate("IMG01", "画像provider", meta["background"]["provider"].endswith("placeholder"),
     meta["background"]["provider"], "Magnific（本番）", stub=True)

# ── 事実系（免除不可。claims.json / evidence.json に依存）──
claims = json.loads((HERE / "claims.json").read_text(encoding="utf-8"))
ev_path = HERE / "evidence.json"
if ev_path.exists():
    ev = {e["claim_id"]: e for e in json.loads(ev_path.read_text(encoding="utf-8"))}
    # 設計書§12 Claim Audit: Evidenceが存在しない事実が1つでもあれば UNSUPPORTED_CLAIM
    unsupported = [c["claim_id"] for c in claims
                   if ev.get(c["claim_id"], {}).get("status") not in ("VERIFIED", "PARTIAL")]
    conflicts = [c["claim_id"] for c in claims if ev.get(c["claim_id"], {}).get("status") == "CONFLICT"]
    # 設計書§11 Double Fact Check: 数字・日付・予測は一次1件 or 一次1+独立1 が必要。
    # 単一の二次情報のみ（PARTIAL）のものは人間レビュー必須項目として明示する。
    HARD = {"number", "date", "forecast"}
    weak = [c["claim_id"] for c in claims
            if c["claim_type"] in HARD and ev.get(c["claim_id"], {}).get("status") == "PARTIAL"]
    verified = sum(1 for c in claims if ev.get(c["claim_id"], {}).get("status") == "VERIFIED")
    gate("FC01", "全claimにEvidence", not unsupported,
         f"未裏付け{len(unsupported)}件 {unsupported[:6]}" if unsupported
         else f"{len(claims)}件（VERIFIED {verified} / PARTIAL {len(claims)-verified}）", "未裏付け0件")
    gate("FC02", "ソース矛盾なし", not conflicts, f"{len(conflicts)}件", "0件")
    gate("FC03", "Double Check充足", not weak,
         f"要確認{len(weak)}件 {weak}" if weak else "全件充足",
         "数値・日付は複数ソース", review=True)
else:
    gate("FC01", "全claimにEvidence", False, "evidence.json 未生成", "0件")
    gate("FC02", "ソース矛盾なし", False, "evidence.json 未生成", "0件")

# ── 出力 ──
print("=" * 92)
print(f"  Final Hard Gates — {len(results)}項目")
print("=" * 92)
blocked, reviews = [], []
for rid, label, ok, actual, expect, stub, review in results:
    if ok and stub:
        mark = "STUB "
    elif ok:
        mark = "PASS "
    elif review:
        mark = "REVIEW"
        reviews.append(rid)
    else:
        mark = "BLOCK"
        blocked.append(rid)
    print(f"  [{mark:<5}] {rid:<8} {label:<22} 実測={actual:<44} 期待={expect}")
print("=" * 92)
if blocked:
    print(f"  結果: BLOCKED — {len(blocked)}項目が不合格 {blocked}")
    print("  → READY_FOR_HUMAN_REVIEW へ進めない（設計書§32）")
else:
    print("  結果: Hard Gates 全項目クリア → READY_FOR_HUMAN_REVIEW へ進める")
    if reviews:
        print(f"  ただし人間レビュー必須項目が {len(reviews)}件: {reviews}")
        print("  → 承認前に必ず内容を確認すること（自動承認は禁止・設計書§42-24）")
sys.exit(1 if blocked else 0)
